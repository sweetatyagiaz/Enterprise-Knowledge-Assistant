from io import BytesIO

import pytest
from docx import Document
from pypdf import PdfWriter

from ai_core.ingestion.loaders import load_document


def test_load_txt():
    assert load_document("notes.txt", b"Hello RAG") == "Hello RAG"


def test_load_txt_with_utf8_bom():
    assert load_document("notes.txt", b"\xef\xbb\xbfHello") == "Hello"


def test_load_markdown():
    content = b"# Enterprise Knowledge Assistant"
    assert load_document("readme.md", content) == "# Enterprise Knowledge Assistant"


def test_load_rejects_unsupported_extension():
    with pytest.raises(ValueError, match="Unsupported file type"):
        load_document("notes.csv", b"a,b")


def test_load_rejects_empty_content():
    with pytest.raises(ValueError, match="document is empty"):
        load_document("notes.txt", b"")


def test_load_rejects_invalid_utf8():
    with pytest.raises(UnicodeDecodeError):
        load_document("notes.txt", b"\xff\xfe\xfa")


def test_load_docx():
    document = Document()
    document.add_paragraph("Enterprise RAG")
    buffer = BytesIO()
    document.save(buffer)

    assert load_document("notes.docx", buffer.getvalue()) == "Enterprise RAG"


def test_load_docx_without_text():
    document = Document()
    buffer = BytesIO()
    document.save(buffer)

    assert load_document("empty.docx", buffer.getvalue()) == ""


def test_load_pdf_without_pages():
    writer = PdfWriter()
    buffer = BytesIO()
    writer.write(buffer)

    assert load_document("empty.pdf", buffer.getvalue()) == ""


def test_load_malformed_pdf():
    with pytest.raises(Exception):
        load_document("broken.pdf", b"not a valid PDF")
